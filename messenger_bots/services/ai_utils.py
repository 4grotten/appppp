"""
AI Assistant utilities for Telegram/WhatsApp bots.

This module consolidates functionality from the external ai_assistant service
to ensure compatibility with the Telegram bot's product parsing.

Prompts can be configured via AIPromptSettings model in Django Admin.
"""
import io
import json
import logging
from typing import Dict, List, Optional, Any

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from django.conf import settings
from django.core.cache import cache

logger = logging.getLogger(__name__)

# Cache key for prompt settings
PROMPT_SETTINGS_CACHE_KEY = "ai_prompt_settings"
PROMPT_SETTINGS_CACHE_TIMEOUT = 300  # 5 minutes


def get_prompt_settings() -> Dict[str, Any]:
    """
    Get AI prompt settings from AIPromptSettings model.
    Uses cache to avoid DB queries on every request.

    Returns empty dict if settings are not active or not available.
    """
    # Try cache first
    cached = cache.get(PROMPT_SETTINGS_CACHE_KEY)
    if cached is not None:
        return cached

    try:
        from common.models import AIPromptSettings
        prompt_data = AIPromptSettings.get_prompt_data()

        # Cache the result
        cache.set(PROMPT_SETTINGS_CACHE_KEY, prompt_data, timeout=PROMPT_SETTINGS_CACHE_TIMEOUT)

        if prompt_data:
            logger.debug(f"[AI_UTILS] Loaded prompt settings from DB (is_active=True)")
        else:
            logger.debug(f"[AI_UTILS] Prompt settings not active, using defaults")

        return prompt_data
    except Exception as e:
        logger.warning(f"[AI_UTILS] Failed to load AIPromptSettings: {e}, using defaults")
        cache.set(PROMPT_SETTINGS_CACHE_KEY, {}, timeout=PROMPT_SETTINGS_CACHE_TIMEOUT)
        return {}


# ============== HTTP Session with Retry ==============

def get_http_session_with_retry(
    retries: int = 3,
    backoff_factor: float = 0.5,
    status_forcelist: tuple = (500, 502, 503, 504),
) -> requests.Session:
    """
    Create HTTP session with automatic retry on failures.

    Args:
        retries: Number of retry attempts
        backoff_factor: Delay multiplier between retries (0.5 = 0.5s, 1s, 2s...)
        status_forcelist: HTTP status codes to retry on

    Returns:
        Configured requests.Session with retry adapter
    """
    session = requests.Session()
    retry_strategy = Retry(
        total=retries,
        backoff_factor=backoff_factor,
        status_forcelist=status_forcelist,
        allowed_methods=["GET", "POST"],
    )
    adapter = HTTPAdapter(max_retries=retry_strategy)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    return session


# ============== File Reading Utilities ==============

def extract_text_from_pdf(file_content: bytes) -> str:
    """Extract text from PDF file content."""
    try:
        # Try pypdf first (newer library)
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(file_content))
            text = ""
            for page in reader.pages:
                text += page.extract_text() or ""
            return text.strip()
        except ImportError:
            pass

        # Fallback to PyPDF2
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(io.BytesIO(file_content))
            text = ""
            for page in reader.pages:
                text += page.extract_text() or ""
            return text.strip()
        except ImportError:
            pass

        # Fallback to pdfplumber
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(file_content)) as pdf:
                text = ""
                for page in pdf.pages:
                    text += page.extract_text() or ""
                return text.strip()
        except ImportError:
            pass

        print("No PDF library available (pypdf, PyPDF2, or pdfplumber)")
        return ""

    except Exception as e:
        print(f"Error extracting text from PDF: {e}")
        return ""


def extract_text_from_docx(file_content: bytes) -> str:
    """Extract text from DOCX file content."""
    try:
        try:
            from docx import Document
            doc = Document(io.BytesIO(file_content))
            return "\n".join([para.text for para in doc.paragraphs]).strip()
        except ImportError:
            pass

        try:
            import docx
            doc = docx.Document(io.BytesIO(file_content))
            return "\n".join([p.text for p in doc.paragraphs]).strip()
        except ImportError:
            pass

        print("No DOCX library available (python-docx)")
        return ""

    except Exception as e:
        print(f"Error extracting text from DOCX: {e}")
        return ""


def format_catalog_json(json_content) -> str:
    """Format catalog JSON into readable text for AI prompt."""
    try:
        if isinstance(json_content, bytes):
            json_content = json_content.decode('utf-8')

        items = json.loads(json_content)
        if not items:
            return "Catalog is empty."

        # Use neutral DATA: format to prevent AI from copying the format
        text = "FULL ORGANIZATION CATALOG:\n"
        text += "(Use scenario_B format with ###NEXT### separator when showing products!)\n\n"

        for item in items:
            name = item.get('name', 'Unknown Item')
            category = item.get('category', 'General')
            price = f"{item.get('price')} {item.get('currency', '')}" if item.get('price') else "Price not set"
            url = item.get('url', 'No link')

            # Use DATA: prefix - raw data, not response format
            text += f"DATA: name={name} | category={category} | price={price} | url={url}\n"

        return text

    except Exception as e:
        logger.warning(f"[CATALOG] Error parsing catalog JSON: {e}")
        return ""


def read_file_from_url(file_url: str) -> str:
    """
    Read and extract text content from a file URL.
    Supports: PDF, DOCX, JSON, TXT, CSV, MD

    Uses HTTP session with automatic retry on server errors.
    """
    try:
        session = get_http_session_with_retry()
        response = session.get(file_url, timeout=(5, 15))  # (connect, read)
        response.raise_for_status()
        file_content = response.content

        # Extract file path without query parameters for extension check
        from urllib.parse import urlparse
        parsed_url = urlparse(file_url)
        file_path = parsed_url.path.lower()

        logger.debug(f"[READ_FILE] URL: {file_url[:100]}...")
        logger.info(f"[READ_FILE] File path={file_path}, size={len(file_content)} bytes")

        if file_path.endswith(".pdf") or ".pdf" in file_url.lower():
            logger.info("[READ_FILE] Detected PDF file")
            return extract_text_from_pdf(file_content)
        elif file_path.endswith(".docx") or ".docx" in file_url.lower():
            logger.info("[READ_FILE] Detected DOCX file")
            return extract_text_from_docx(file_content)
        elif file_path.endswith(".json") or ".json" in file_url.lower():
            return format_catalog_json(file_content)
        elif file_path.endswith((".txt", ".csv", ".log", ".md")):
            try:
                return file_content.decode("utf-8")[:5000]
            except UnicodeDecodeError:
                logger.warning(f"[READ_FILE] Failed to decode file as UTF-8: {file_url}")
                return ""
        else:
            # Try to detect by content type or magic bytes
            logger.info("[READ_FILE] Unknown extension, trying to detect type...")
            # PDF magic bytes: %PDF
            if file_content[:4] == b'%PDF':
                logger.info("[READ_FILE] Detected PDF by magic bytes")
                return extract_text_from_pdf(file_content)
            # DOCX is a ZIP file starting with PK
            elif file_content[:2] == b'PK':
                logger.info("[READ_FILE] Detected DOCX/ZIP by magic bytes")
                return extract_text_from_docx(file_content)
            else:
                logger.warning(f"[READ_FILE] Unsupported file type: {file_path}")
                return ""

    except requests.exceptions.RequestException as e:
        logger.warning(f"[READ_FILE] Error reading file from URL: {e}")
        return ""


# ============== Prompt Building ==============

# Default prompts (used when AIPromptSettings is not active)
DEFAULT_PROMPTS = {
    "language_detection_rule": (
        "🌐 CRITICAL LANGUAGE RULE:\n"
        "Detect the language from USER'S MESSAGES (not from any settings).\n"
        "- If user writes in English (Hello, What can you do, etc.) → respond in ENGLISH\n"
        "- If user writes in Russian (Привет, Что умеешь, etc.) → respond in RUSSIAN\n"
        "- If user explicitly asks 'Speak English' or 'Говори по-русски' → switch to that language\n"
        "- Translate all data (products, contacts) to the user's language."
    ),
    "formatting_rules": (
        "⛔ STRICT FORMATTING RULES:\n"
        "1. NO MARKDOWN. No *, **, [text](url).\n"
        "2. Send LINKS as plain text only.\n"
        "3. SEPARATOR: Use '###NEXT###' to separate different products or the final link."
    ),
    "scenario_a_discounts": (
        "scenario_A: DISCOUNTS & COUPONS\n"
        "   - IF user asks about discounts, coupons, or bonuses:\n"
        "   - Answer ONLY about the promotions.\n"
        "   - DO NOT list products/items unless the user explicitly asks for them.\n"
        "   - DO NOT use the ###NEXT### tag in this scenario."
    ),
    "scenario_b_products": (
        "scenario_B: PRODUCTS (Catalogue)\n"
        "   ⚠️ CRITICAL: You MUST use this EXACT format for EACH product:\n"
        "   - DO NOT use dashes (-) or bullet points!\n"
        "   - DO NOT copy the DATA: format from catalog!\n"
        "   - Put ###NEXT### BETWEEN each product (not at the end)\n\n"
        "   FORMAT for Russian:\n"
        "   Товар: <item name>\n"
        "   Цена: <price>\n"
        "   Ссылка: <url>\n"
        "   ###NEXT###\n"
        "   FORMAT for English:\n"
        "   Product: <item name>\n"
        "   Price: <price>\n"
        "   Link: <url>"
    ),
    "scenario_c_contacts": (
        "scenario_C: CONTACTS\n"
        "   - IF user asks for contacts/address/phone:\n"
        "   - 1. First check the 'KNOWLEDGE BASE' (files/answers) below.\n"
        "   - 2. If not found, use 'ORGANIZATION DATA' below.\n"
        "   - Format for Russian: 📞 Телефон: / 🏢 Адрес: / 🕘 Часы работы:\n"
        "   - Format for English: 📞 Phone: / 🏢 Address: / 🕘 Hours:"
    ),
    "scenario_d_general": (
        "scenario_D: GENERAL QUESTIONS\n"
        "   - IF user asks general questions (Привет, Hello, etc.):\n"
        "   - Answer naturally and helpfully.\n"
        "   - Briefly describe what you can help with (products, promotions, contacts).\n"
        "   - DO NOT use ###NEXT### tag.\n"
        "   - DO NOT list products unless asked."
    ),
    "ending_rule": (
        "🏁 ENDING RULE:\n"
        "   - ONLY when listing products, finish with organization link.\n"
        "   - Russian: ###NEXT###\nБольше товаров на странице: {org_page_url}\n"
        "   - English: ###NEXT###\nMore items at: {org_page_url}"
    ),
    "search_rules": (
        "SEARCH RULES:\n"
        "- Extract keywords from user question (e.g. 'купальник', 'sneakers', 'кроссовки', 'dress')\n"
        "- Search ENTIRE catalog for items matching keywords in name/category/description\n"
        "- If found - show ALL matching products, not just first ones\n"
        "- If not found - say so and suggest similar categories (in user's language)"
    ),
    "few_shot_example_greeting_ru": "Здравствуйте! Я помощник {organization}. Могу помочь с информацией о товарах, акциях и контактах. Чем могу быть полезен?",
    "few_shot_example_greeting_en": "Hello! I'm an assistant at {organization}. I can help with product info, promotions, and contacts. How can I help you?",
    "few_shot_example_capabilities_ru": "Я могу помочь вам с информацией о товарах в {organization}, рассказать об акциях и скидках, предоставить контактные данные и адрес. Задавайте вопросы!",
    "few_shot_example_capabilities_en": "I can help you with product information at {organization}, tell you about promotions and discounts, provide contact details and address. Feel free to ask!",
    "few_shot_example_contacts": "📞 Телефон: +7 XXX XXX-XX-XX\n🏢 Адрес: ул. Примерная, 1\n🕘 Часы работы: 10:00 - 20:00",
}


def build_system_prompt(
    assistant_info: Dict[str, Any],
    organization_info: Dict[str, Any],
    organization_page_url: str,
    qa_pairs: List[Dict[str, Any]],
    catalog_content: str = "",
    marketing_info: List[str] = None,
    item_info: Dict[str, Any] = None,
    user_language: str = "ru",  # kept for compatibility but not used for forcing language
    cached_file_contents: Dict[str, str] = None,
) -> str:
    """
    Build a comprehensive system prompt for the AI assistant.

    Uses prompts from AIPromptSettings if active, otherwise falls back to defaults.

    The product format is designed to be compatible with telegram.py's
    parse_products_from_response() function which expects:
    - "Товар:" or "Product:" at the start
    - "Ссылка:" or "Link:" with URL at the end

    Language is detected from message context, not from Telegram settings.
    """
    # Get prompt settings from DB (cached)
    prompt_settings = get_prompt_settings()

    # Helper to get setting with fallback to default
    def get_setting(key: str) -> str:
        return prompt_settings.get(key) or DEFAULT_PROMPTS.get(key, "")

    # Extract contact info
    phones = organization_info.get('phones', [])
    if isinstance(phones, list):
        phones_str = ", ".join([str(p) for p in phones])
    else:
        phones_str = str(phones)
    address_str = organization_info.get('address', 'Unknown')
    opens_at = organization_info.get('opens_at', 'Unknown')
    closes_at = organization_info.get('closes_at', 'Unknown')
    social_links = organization_info.get('social_links', '')

    # Marketing info
    marketing_str = ""
    if marketing_info:
        marketing_str = "\n".join([f"- {m}" for m in marketing_info])

    # Build identity from template or defaults
    identity_template = prompt_settings.get('identity_template') or (
        "You are {assistant_name}, an assistant at {organization}.\n"
        "Position: {position}. Gender: {gender}."
    )
    identity = identity_template.format(
        assistant_name=assistant_info.get('name', 'Assistant'),
        organization=assistant_info.get('organization', 'organization'),
        position=assistant_info.get('position', 'consultant'),
        gender=assistant_info.get('gender', 'not specified'),
    )

    # Build ending rule with org URL
    ending_rule = get_setting('ending_rule').format(org_page_url=organization_page_url)

    # Build prompt using settings from DB or defaults
    prompt = (
        f"{get_setting('language_detection_rule')}\n\n"
        f"IDENTITY:\n{identity}\n\n"
        f"{get_setting('formatting_rules')}\n\n"
        "🧠 LOGIC SCENARIOS:\n\n"
        f"{get_setting('scenario_a_discounts')}\n\n"
        f"{get_setting('scenario_b_products')}\n\n"
        f"{get_setting('scenario_c_contacts')}\n"
        f"   - Socials: {social_links} (if available)\n\n"
        f"{get_setting('scenario_d_general')}\n\n"
        f"{ending_rule}\n\n"
        "=== DATA SECTIONS ===\n\n"
    )

    # Add marketing/promotions
    if marketing_str:
        prompt += f"💰 ACTIVE PROMOTIONS:\n{marketing_str}\n\n"

    # Add organization data
    prompt += (
        f"🏢 ORGANIZATION DATA (Backup for contacts):\n"
        f"Phone: {phones_str}\n"
        f"Address: {address_str}\n"
        f"Open: {opens_at}\n"
        f"Close: {closes_at}\n\n"
    )

    # Add current item info if user is viewing a specific product
    if item_info:
        prompt += (
            f"📦 USER IS LOOKING AT THIS ITEM:\n"
            f"ID: {item_info.get('id')}\n"
            f"Name: {item_info.get('name')}\n"
            f"Price: {item_info.get('price')}\n"
            f"Description: {item_info.get('description')}\n\n"
        )

    # Add Q&A training data (KNOWLEDGE BASE)
    logger.info(f"[BUILD_PROMPT] qa_pairs_count={len(qa_pairs) if qa_pairs else 0}")
    if qa_pairs:
        prompt += "📚 KNOWLEDGE BASE (Primary source for specific questions):\n"
        prompt += "IMPORTANT: When user asks for a link/file mentioned in answers below, provide the File URL!\n\n"
        for qa in qa_pairs:
            question = qa.get('question') or ''
            answer = qa.get('answer') or ''
            logger.debug(f"[BUILD_PROMPT] Adding Q&A: Q='{question[:50]}...' A='{answer[:50]}...'")
            if question and answer:
                prompt += f"Q: {question}\nA: {answer}\n"

            # Add file contents and URLs - support both new and legacy keys
            files_to_read = qa.get('files_to_read') or qa.get('files') or []
            files_to_send = qa.get('files_to_send') or []
            logger.debug(
                f"[BUILD_PROMPT] Q&A files: readable={len(files_to_read)}, send_only={len(files_to_send)}"
            )

            for file_url in files_to_read:
                if file_url:
                    # Add readable files: AI can share URL and use content
                    prompt += f"📎 File URL (share this when asked): {file_url}\n"

                    # Use cached content if available, otherwise download fresh
                    if cached_file_contents and file_url in cached_file_contents:
                        file_content = cached_file_contents[file_url]
                        logger.debug(f"[BUILD_PROMPT] Using cached file content: {file_url[:50]}...")
                    else:
                        logger.info(f"[BUILD_PROMPT] Loading file (no cache): {file_url}")
                        file_content = read_file_from_url(file_url)

                    if file_content:
                        prompt += f"File content preview: {file_content[:1500]}\n"
                    else:
                        logger.warning("[BUILD_PROMPT] Could not read file content, but URL is added")

            # Add non-readable files as links only
            for file_url in files_to_send:
                if file_url:
                    prompt += f"📎 File URL (share this when asked): {file_url}\n"
            prompt += "\n"
    else:
        logger.warning("[BUILD_PROMPT] No Q&A pairs provided")

    # Add catalog with search rules
    if catalog_content:
        prompt += f"🛒 PRODUCT CATALOG:\n{catalog_content}\n"
        prompt += f"{get_setting('search_rules')}\n"
        prompt += "- Format each product using scenario_B format above\n\n"

    # Add few-shot examples for BOTH languages (from settings or defaults)
    org_name = assistant_info.get('organization', 'organization')

    # Get examples from settings and format with org name
    greeting_ru = get_setting('few_shot_example_greeting_ru').format(organization=org_name)
    greeting_en = get_setting('few_shot_example_greeting_en').format(organization=org_name)
    capabilities_ru = get_setting('few_shot_example_capabilities_ru').format(organization=org_name)
    capabilities_en = get_setting('few_shot_example_capabilities_en').format(organization=org_name)
    contacts_example = get_setting('few_shot_example_contacts')

    prompt += (
        "📝 RESPONSE EXAMPLES (use language matching user's messages):\n\n"
        "--- RUSSIAN EXAMPLES ---\n"
        f"User: Привет!\nAssistant: {greeting_ru}\n\n"
        f"User: Что ты умеешь?\nAssistant: {capabilities_ru}\n\n"
        f"User: Контакты\nAssistant: {contacts_example}\n\n"
        "--- ENGLISH EXAMPLES ---\n"
        f"User: Hello!\nAssistant: {greeting_en}\n\n"
        f"User: What can you do?\nAssistant: {capabilities_en}\n\n"
        "User: Contacts\nAssistant: 📞 Phone: +7 XXX XXX-XX-XX\n🏢 Address: Example St. 1\n🕘 Working hours: 10:00 - 20:00\n\n"
    )

    # Final reminder about language detection
    prompt += (
        "\n⚠️ REMINDER: Always detect language from user's CURRENT message. "
        "If user switches language mid-conversation, switch your response language too.\n"
    )

    print(f"[BUILD_PROMPT] Final prompt length: {len(prompt)} chars")
    return prompt


# ============== OpenAI API ==============

def call_openai(
    question: str,
    system_prompt: str,
    chat_history: Optional[List[Dict[str, str]]] = None,
    model: str = "gpt-4o-mini",  # 128k context, cheaper than gpt-3.5-turbo
    max_tokens: int = 1500,
    temperature: float = 0.5,
) -> str:
    """
    Call OpenAI API via ai_assistant proxy service.

    The proxy is needed because the backend server may be in a region
    where OpenAI is blocked, but ai_assistant server is in allowed region.

    Args:
        question: User's question
        system_prompt: System prompt with training data
        chat_history: Previous messages for context
        model: OpenAI model name
        max_tokens: Maximum response tokens
        temperature: Response creativity (0-1)

    Returns:
        AI response text
    """
    # Get AI Assistant service URL for proxy
    ai_assistant_url = getattr(settings, "AI_ASSISTANT_URL", "http://161.35.153.151:8080")
    proxy_url = f"{ai_assistant_url}/bot/openai-proxy/"

    try:
        payload = {
            "system_prompt": system_prompt,
            "question": question,
            "chat_history": chat_history or [],
            "model": model,
            "max_tokens": max_tokens,
            "temperature": temperature,
        }

        print(f"[AI_UTILS] Calling OpenAI proxy at {proxy_url}")
        logger.info(
            "[AI_PROXY_MODEL_TRACE] Sending request with model=%s max_tokens=%s chat_history_items=%s",
            model,
            max_tokens,
            len(chat_history) if isinstance(chat_history, list) else 0,
        )
        session = get_http_session_with_retry()
        response = session.post(
            proxy_url,
            json=payload,
            timeout=(5, 30)  # (connect_timeout, read_timeout) - fast connection, reasonable read
        )

        result = response.json()

        if "answer" in result:
            logger.info(
                "[AI_PROXY_MODEL_TRACE] Response received for model=%s status=%s",
                model,
                response.status_code,
            )
            answer = result["answer"]
            print(f"[AI_UTILS] OpenAI proxy response received: {len(answer)} chars")
            return answer
        elif "error" in result:
            print(f"[AI_UTILS] OpenAI proxy error: {result['error']}")
            return ""
        else:
            print(f"[AI_UTILS] Unexpected proxy response format: {result}")
            return ""

    except Exception as e:
        print(f"[AI_UTILS] Error calling OpenAI proxy: {e}")
        return ""


def get_openai_config():
    """Get OpenAI configuration from settings or cache."""
    from django.core.cache import cache

    api_key = cache.get('openai_api_key')
    model_name = cache.get('openai_model_name')

    if not api_key:
        api_key = getattr(settings, 'OPENAI_API_KEY', '')

    if not model_name:
        model_name = getattr(settings, 'OPENAI_MODEL_NAME', 'gpt-3.5-turbo')

    return api_key, model_name
